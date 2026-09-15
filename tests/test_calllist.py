"""The call list: ranking, the plain-English reason, and the ICP filter.

A score on its own tells whoever is dialling nothing. These tests are mostly
about the `Why` line — that it says only what was actually collected, and never
pads itself out with things nobody verified.
"""

import pytest

import storage
from storage import explain_lead, is_icp, LeadStore


def lead(**overrides):
    base = {"NPI Number": "1234567890", "Provider First Name": "Jane",
            "Provider Last Name": "Smith", "State": "TX",
            "Source": "Psychology Today"}
    base.update(overrides)
    return base


class TestExplanation:

    def test_an_owner_is_named_as_one(self):
        assert "Owner" in explain_lead({"Role": "Owner", "Phone": "x"})

    def test_a_non_owner_role_is_qualified(self):
        """Clinical Director reads like a decision maker and is not one."""
        why = explain_lead({"Role": "Clinical Director"})
        assert "not the owner" in why

    def test_practice_size_is_described_in_words(self):
        assert "solo practice" in explain_lead({"Practice Size": 1})
        assert "small practice (14 clinicians)" in explain_lead({"Practice Size": 14})
        assert "large group (676 clinicians)" in explain_lead({"Practice Size": 676})

    def test_billing_is_shortened_for_reading(self):
        assert "$231k Medicare billing" in explain_lead({"Annual Billing": 230570})
        assert "$2.4M Medicare billing" in explain_lead({"Annual Billing": 2400000})

    def test_revenue_per_clinician_is_flagged_for_small_practices(self):
        """A lot of money through very few people means no billing department —
        the sharpest RCM signal available."""
        why = explain_lead({"Annual Billing": 1400000, "Practice Size": 4})
        assert "/clinician" in why

    def test_revenue_per_clinician_is_not_flagged_for_large_groups(self):
        why = explain_lead({"Annual Billing": 1400000, "Practice Size": 900})
        assert "/clinician" not in why

    def test_phone_quality_is_stated(self):
        assert "direct line" in explain_lead(
            {"Phone Source": "Provider Website (personal/direct (tel: link))"})
        assert "shared front desk" in explain_lead(
            {"Phone Source": "Provider Website (shared practice line)"})
        assert "directory number only" in explain_lead(
            {"Phone Source": "Psychology Today (listing)"})

    def test_a_federally_confirmed_phone_is_called_out(self):
        why = explain_lead({"Phone Source": "", "Phone Verified": storage.PHONE_CONFIRMED})
        assert "confirmed federally" in why

    def test_email_quality_is_stated(self):
        assert "no email yet" in explain_lead({"Phone": "(512) 471-3300"})
        assert "email confirmed" in explain_lead(
            {"Email": "a@b.example", "Email Status": "valid"})
        assert "shared mailbox" in explain_lead(
            {"Email": "info@b.example", "Email Status": "role"})
        assert "someone else's" in explain_lead(
            {"Email": "x@gmail.com", "Email Status": "risky"})

    def test_nothing_is_invented_for_an_empty_lead(self):
        """The line must never imply data that was never collected."""
        why = explain_lead({})
        assert why == "contact details only"
        for invented in ("solo", "billing", "owner", "confirmed"):
            assert invented not in why.lower()

    def test_a_rich_lead_reads_as_a_sentence(self):
        why = explain_lead({
            "Role": "Owner", "Practice Size": 4, "Annual Billing": 620000,
            "Phone Source": "Provider Website (personal/direct)",
            "Email": "jane@practice.example", "Email Status": "valid"})
        assert why.startswith("Owner")
        assert "small practice (4 clinicians)" in why
        assert "$620k Medicare billing" in why
        assert "direct line" in why
        assert "email confirmed" in why


class TestICPFilter:

    def test_an_owner_of_a_small_practice_qualifies(self):
        assert is_icp({"Role": "Owner", "Practice Size": 6, "Email": "a@b.example"})

    def test_a_non_owner_never_qualifies(self):
        assert not is_icp({"Role": "Clinical Director", "Practice Size": 6,
                           "Email": "a@b.example"})

    def test_a_large_group_is_excluded(self):
        """Above ~25 clinicians a practice has its own billing department."""
        assert not is_icp({"Role": "Owner", "Practice Size": 900,
                           "Email": "a@b.example"})

    def test_unknown_practice_size_does_not_disqualify(self):
        """Most therapists will never have CMS data; excluding them would throw
        away a whole source."""
        assert is_icp({"Role": "Owner", "Email": "a@b.example"})

    def test_an_unreachable_lead_is_excluded(self):
        assert not is_icp({"Role": "Owner", "Practice Size": 4})


class TestBestLeads:

    def test_leads_come_back_best_first(self, store):
        store.add_leads_batch([
            lead(**{"NPI Number": "1111111111", "Phone": "(512) 471-3300"}),
            lead(**{"NPI Number": "2222222222", "Role": "Owner",
                    "Email": "a@b.example", "Email Status": "valid",
                    "Phone": "(512) 471-3399",
                    "Phone Source": "Provider Website (personal/direct)"}),
        ])
        best = store.best_leads(limit=10)
        assert best[0]["NPI Number"] == "2222222222"
        assert int(best[0]["Lead Score"]) > int(best[1]["Lead Score"])

    def test_every_row_carries_a_reason(self, store):
        store.add_leads_batch([lead(Phone="(512) 471-3300", Role="Owner")])
        assert store.best_leads(limit=5)[0]["Why"]

    def test_unreachable_leads_are_left_out(self, store):
        store.add_leads_batch([lead(**{"NPI Number": "3333333333"})])
        assert store.best_leads(limit=10) == []

    def test_the_limit_is_respected(self, store):
        store.add_leads_batch([
            lead(**{"NPI Number": "111111111{0}".format(i), "Phone": "(512) 471-330{0}".format(i)})
            for i in range(5)])
        assert len(store.best_leads(limit=3)) == 3

    def test_owners_only_filters(self, store):
        store.add_leads_batch([
            lead(**{"NPI Number": "1111111111", "Phone": "(512) 471-3300"}),
            lead(**{"NPI Number": "2222222222", "Phone": "(512) 471-3399",
                    "Role": "Owner"}),
        ])
        best = store.best_leads(limit=10, owners_only=True)
        assert len(best) == 1
        assert best[0]["Role"] == "Owner"

    def test_state_filtering(self, store):
        store.add_leads_batch([
            lead(**{"NPI Number": "1111111111", "Phone": "(512) 471-3300"}),
            lead(**{"NPI Number": "2222222222", "Phone": "(614) 545-0123",
                    "State": "OH"}),
        ])
        assert len(store.best_leads(limit=10, states=["OH"])) == 1


class TestCallListSheet:

    def test_the_call_list_is_the_first_tab(self, store, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        from lead_manager import ExcelExporter

        store.add_leads_batch([
            lead(**{"NPI Number": "111111111{0}".format(i),
                    "Phone": "(512) 471-330{0}".format(i), "Role": "Owner"})
            for i in range(3)])

        path = tmp_path / "out.xlsx"
        ExcelExporter(str(path)).export(store, call_list=50)
        workbook = openpyxl.load_workbook(path)
        try:
            assert "Call List" in workbook.sheetnames
            sheet = workbook["Call List"]
            headers = [c.value for c in sheet[1]]
            assert headers[:3] == ["#", "Name", "Practice"]
            assert "Why this lead" in headers
            # Blank columns for the caller to fill in as they work the list.
            assert "Outcome" in headers and "Notes" in headers
            assert sheet.max_row == 4  # header + 3 leads
        finally:
            workbook.close()

    def test_the_call_list_can_be_omitted(self, store, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        from lead_manager import ExcelExporter

        store.add_leads_batch([lead(Phone="(512) 471-3300")])
        path = tmp_path / "out.xlsx"
        ExcelExporter(str(path)).export(store, call_list=0)
        workbook = openpyxl.load_workbook(path)
        try:
            assert "Call List" not in workbook.sheetnames
        finally:
            workbook.close()

    def test_scores_are_written_as_numbers(self, store, tmp_path):
        """Or Excel sorts the column as text and puts 9 above 80."""
        pytest.importorskip("openpyxl")
        import openpyxl
        from lead_manager import ExcelExporter

        store.add_leads_batch([lead(Phone="(512) 471-3300", Role="Owner")])
        path = tmp_path / "out.xlsx"
        ExcelExporter(str(path)).export(store, call_list=10)
        workbook = openpyxl.load_workbook(path)
        try:
            sheet = workbook["Call List"]
            score = sheet.cell(row=2, column=[c.value for c in sheet[1]].index("Score") + 1).value
            assert isinstance(score, (int, float))
        finally:
            workbook.close()
