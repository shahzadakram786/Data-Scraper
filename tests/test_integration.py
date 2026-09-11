"""End-to-end: collect, enrich, verify, qualify, send, export.

Every other test file checks one component. This one checks that they still fit
together — that a lead collected in step one carries its phone trust level,
verification status, CMS figures and score all the way through to the Call List
tab and the outbound message. Unit tests pass happily while the seams rot.

Everything external is faked. The point is the wiring, not the internet.
"""

import pytest

import storage
import outreach
from storage import LeadStore, RUN_OK
from pipeline import LeadPipeline, Reporter
from lead_manager import ExcelExporter


# ------------------------------------------------------------------- fakes

class FakePT:
    """Two providers: one owner with a website, one bare listing."""

    PROVIDERS = [
        {"pt_id": "1111111", "profile_url": "https://pt.example/1111111",
         "name": "Jane Smith", "credential": "Psychologist, PsyD",
         "phone": "(512) 471-3300", "city": "Austin", "state_full": "TX",
         "zip": "78701", "is_org": False},
        {"pt_id": "2222222", "profile_url": "https://pt.example/2222222",
         "name": "Bob Jones", "credential": "Counselor, LPC",
         "phone": "(214) 692-0177", "city": "Dallas", "state_full": "TX",
         "zip": "75201", "is_org": False},
    ]

    def iter_listing(self, state, stop_event=None, max_pages=None, start_page=1):
        yield 1, list(self.PROVIDERS)

    def _resolve_website(self, pt_id):
        return "https://smithpsych.example" if pt_id == "1111111" else ""

    def get_profile_details(self, pt_id, url):
        return {"specialty": "", "website": self._resolve_website(pt_id)}


class FakeEmailEngine:
    def scrape_website_contact_info(self, url, person_name=None):
        # The owner's site yields a direct line, a role and an address.
        return {"email": "jane@smithpsych.example",
                "phone": "(512) 471-3399",
                "phone_confidence": "personal/direct (tel: link)",
                "role": "Owner",
                "socials": {"LinkedIn": "https://linkedin.com/in/jane"}}

    def find_email_by_phone(self, phone, person_name=None, org_name=None):
        return ""

    def find_email(self, lead):
        # The second pass finds one for the provider who had no website.
        lead["Email"] = "bob@jonescounseling.example"
        lead["Website"] = "https://jonescounseling.example"
        return lead

    def find_all_emails_by_phone(self, phone, **kwargs):
        return []

    def find_all_emails_by_name(self, name, city="", state="", **kwargs):
        return []

    @staticmethod
    def is_valid_phone(phone):
        from email_engine import EmailEngine
        return EmailEngine.is_valid_phone(phone)

    @staticmethod
    def is_tollfree(phone):
        from email_engine import EmailEngine
        return EmailEngine.is_tollfree(phone)

    @staticmethod
    def is_ownership_role(role):
        from email_engine import EmailEngine
        return EmailEngine.is_ownership_role(role)

    @staticmethod
    def is_platform_site(url):
        from email_engine import EmailEngine
        return EmailEngine.is_platform_site(url)

    @staticmethod
    def _registrable_domain(netloc):
        from email_engine import EmailEngine
        return EmailEngine._registrable_domain(netloc)


class FakeNPI:
    def find_by_name(self, first, last, state, city=None):
        if last == "Smith":
            return {"NPI Number": "1234567890", "Address": "1 Main St",
                    "City": "Austin", "Zip": "78701",
                    "Specialty": "Psychologist", "Organization Name": ""}
        return None

    def find_by_organization(self, name, state, city=None):
        return None

    def find_by_number(self, number):
        return None

    def fetch_providers(self, state, specialty, **kwargs):
        return []


class FakeCMS:
    def billing_for(self, npi):
        if npi == "1234567890":
            return {"annual_billing": 480000.0, "medicare_paid": 96000.0,
                    "patient_volume": 410, "service_count": 900}
        return {}

    def enrollment_for(self, npi):
        if npi == "1234567890":
            return {"practice_name": "Smith Psychological Services",
                    "practice_size": 6, "practice_id": "9988776655",
                    "filed_phone": "5124713300", "specialty": "PSYCHOLOGIST",
                    "telehealth": True, "city": "AUSTIN", "state": "TX"}
        return {}


SETTINGS = {
    "from_name": "Shahzad Akram", "from_email": "shahzad@example.com",
    "postal_address": "12 Example Street, Lahore",
    "unsubscribe_email": "unsubscribe@example.com",
    "smtp_host": "smtp.example.com", "smtp_user": "shahzad@example.com",
}


@pytest.fixture
def verifier(offline_verifier):
    instance = offline_verifier()
    instance.mx_by_domain = {
        "smithpsych.example": ["mx.smithpsych.example"],
        "jonescounseling.example": ["mx.jonescounseling.example"],
    }
    instance.existing_mailboxes = {"jane@smithpsych.example",
                                   "bob@jonescounseling.example"}
    return instance


@pytest.fixture
def pipeline(store, verifier):
    return LeadPipeline(
        store=store, npi_engine=FakeNPI(), email_engine=FakeEmailEngine(),
        pt_engine=FakePT(), verifier=verifier, cms_engine=FakeCMS(),
        reporter=Reporter(), settings=SETTINGS)


# --------------------------------------------------------------- the chain

class TestFullChain:

    def test_collect_enrich_verify_qualify_send_export(self, store, pipeline, tmp_path):
        # 1. Collect --------------------------------------------------------
        result = pipeline.collect_psychology_today("Texas", "TX", resume=False)
        assert result.status == RUN_OK
        assert result.added == 2

        jane = self._find(store, "Smith")
        assert jane["NPI Number"] == "1234567890", "NPI cross-reference attached"
        assert jane["Phone"] == "(512) 471-3399", "the direct line beat the listing"
        assert "personal/direct" in jane["Phone Source"]
        assert jane["Role"] == "Owner"
        assert jane["Email Status"] == "valid", "verified during collection"

        bob = self._find(store, "Jones")
        assert bob["NPI Number"] == "PT-2222222", "no registry match, traceable id"
        assert bob["Email"] == "", "no website, so no address yet"
        assert bob["Phone"] == "(214) 692-0177", "fell back to the listing number"

        # 2. Fill in the missing address ------------------------------------
        assert store.count(missing_email=True) == 1
        pipeline.find_missing_emails()
        bob = self._find(store, "Jones")
        assert bob["Email"] == "bob@jonescounseling.example"
        assert bob["Email Status"] == "valid"

        # 3. CMS enrichment --------------------------------------------------
        pipeline.enrich_from_cms()
        jane = self._find(store, "Smith")
        assert float(jane["Annual Billing"]) == 480000.0
        assert int(jane["Practice Size"]) == 6
        assert jane["Phone Verified"] == storage.PHONE_DIFFERS, (
            "the scraped direct line differs from the billing office on file")
        # Bob carries a "PT-" placeholder rather than a real NPI, so there is
        # nothing for CMS to look up and the pass correctly never touches him.
        # A blank here means "not checkable", which is distinct from the
        # "No federal enrollment record" written when a lookup found nothing.
        bob = self._find(store, "Jones")
        assert bob["Phone Verified"] == ""
        assert store.count(needs_cms=True) == 0, "nothing left that can be checked"

        # 4. Scoring and ranking ---------------------------------------------
        best = store.best_leads(limit=10)
        assert best[0]["Provider Last Name"] == "Smith", "the owner ranks first"
        assert int(best[0]["Lead Score"]) > int(best[1]["Lead Score"])
        why = best[0]["Why"]
        for expected in ("Owner", "small practice (6 clinicians)",
                         "$480k Medicare billing", "direct line", "email confirmed"):
            assert expected in why, why

        # 5. Send -------------------------------------------------------------
        sent = []
        original = outreach.EmailSender

        class FakeSender:
            def __init__(self, settings, logger=None):
                pass

            def send(self, to, subject, body):
                sent.append((to, subject, body))
                return True, "sent"

            def close(self):
                pass

        outreach.EmailSender = FakeSender
        try:
            send_result = pipeline.send_campaign(
                "About billing at {{Practice}}",
                "Hi {{First Name}}, I work with practices in {{State}}.",
                campaign="integration", min_seconds=0, max_seconds=0)
        finally:
            outreach.EmailSender = original

        assert send_result.status == RUN_OK
        assert len(sent) == 2
        recipients = {to for to, _s, _b in sent}
        assert recipients == {"jane@smithpsych.example",
                              "bob@jonescounseling.example"}
        # Merge fields resolved, and the CMS practice name reached the message.
        subjects = {s for _t, s, _b in sent}
        assert "About billing at Smith Psychological Services" in subjects
        assert all(body.startswith("Hi ") for _t, _s, body in sent)

        assert store.outreach_stats()["sent"] == 2
        assert self._find(store, "Smith")["Outreach Status"] == "Contacted"

        # Nobody gets a second copy.
        more, excluded = pipeline.recipients_for()
        assert more == []
        assert excluded.get("already contacted") == 2

        # 6. Export ------------------------------------------------------------
        path = tmp_path / "final.xlsx"
        written = ExcelExporter(str(path)).export(store, call_list=50)
        assert written == 2

        import openpyxl
        workbook = openpyxl.load_workbook(path)
        try:
            assert workbook.sheetnames[0] == "Summary"
            assert "Call List" in workbook.sheetnames
            assert "Texas - Psychology Today" in workbook.sheetnames

            call_list = workbook["Call List"]
            headers = [c.value for c in call_list[1]]
            top_row = {headers[i]: c.value
                       for i, c in enumerate(call_list[2]) if i < len(headers)}
            assert top_row["Name"] == "Jane Smith"
            assert top_row["Outcome"] == "Contacted", "the send advanced the status"
            assert isinstance(top_row["Score"], int)
            assert "Owner" in top_row["Why this lead"]

            data = workbook["Texas - Psychology Today"]
            data_headers = [c.value for c in data[1]]
            assert data_headers == storage.HEADERS, "export schema is stable"
        finally:
            workbook.close()

    def test_a_second_collection_adds_nothing_and_says_so(self, store, pipeline):
        """Re-running a finished state must not duplicate it, and must not
        report success either — an empty result is the signature of a broken
        scraper and has to stay visible."""
        pipeline.collect_psychology_today("Texas", "TX", resume=False)
        again = pipeline.collect_psychology_today("Texas", "TX", resume=False)
        assert again.status == storage.RUN_EMPTY
        assert store.get_stats()["total"] == 2

    def test_a_bounce_removes_the_lead_from_all_future_campaigns(self, store, pipeline):
        pipeline.collect_psychology_today("Texas", "TX", resume=False)
        original = outreach.EmailSender

        class BouncingSender:
            def __init__(self, settings, logger=None):
                pass

            def send(self, to, subject, body):
                return False, "rejected: 550 no such user"

            def close(self):
                pass

        outreach.EmailSender = BouncingSender
        try:
            pipeline.send_campaign("Hi", "Body", min_seconds=0, max_seconds=0)
        finally:
            outreach.EmailSender = original

        assert store.is_suppressed("jane@smithpsych.example")
        recipients, excluded = pipeline.recipients_for()
        assert recipients == []
        assert excluded.get("opted out or previously bounced")

    def test_stats_stay_consistent_across_the_whole_chain(self, store, pipeline):
        pipeline.collect_psychology_today("Texas", "TX", resume=False)
        pipeline.find_missing_emails()
        pipeline.enrich_from_cms()

        stats = store.get_stats()
        assert stats["total"] == 2
        assert stats["with_email"] == 2
        assert stats["with_phone"] == 2
        assert stats["verified_email"] == 2
        assert stats["with_billing"] == 1
        assert stats["by_state"]["TX"]["total"] == 2
        assert sum(v["total"] for v in stats["by_source"].values()) == stats["total"]

    @staticmethod
    def _find(store, last_name):
        for lead in store.iter_leads():
            if lead["Provider Last Name"] == last_name:
                return lead
        raise AssertionError("no lead named " + last_name)
